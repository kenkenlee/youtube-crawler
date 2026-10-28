"""Tags stored for botalk-news sync: JSON array on the video, exposed by list/detail/export."""
import asyncio
import csv
import io
from datetime import datetime

from app.models.channel import Channel
from app.models.crawl_session import CrawlSession
from app.models.video import Video
from app.schemas.video import VideoWithChannel
from app.services.export_service import ExportService
from app.services.video_tags import (
    backfill_missing_tags,
    normalize_tags,
    resolve_tags,
    tags_missing,
)
from app.services.youtube_service import YouTubeService


def _settings_get(key, default=None):
    if key in ("AUTO_FETCH_TRANSCRIPT", "AUTO_SUMMARIZE"):
        return False
    return default


def _channel(db, name="Chan"):
    channel = Channel(
        channel_id=f"ref-{name}",
        youtube_channel_id="UCtest",
        channel_name=name,
        channel_url="https://www.youtube.com/channel/UCtest",
        keywords=["news"],
    )
    db.add(channel)
    db.commit()
    db.refresh(channel)
    return channel


def _session(db, channel):
    session = CrawlSession(
        session_name="tags",
        channel_ids=[channel.id],
        status="running",
        max_videos=10,
        total_videos_found=0,
        videos_processed=0,
        error_count=0,
    )
    db.add(session)
    db.commit()
    db.refresh(session)
    return session


def _video(db, channel, video_id, tags=None, checked=False, title="Title"):
    video = Video(
        channel_id=channel.id,
        video_id=video_id,
        title=title,
        description="desc",
        tags=tags,
        tags_checked_at=datetime.utcnow() if checked else None,
        summary_text="A summary",
        created_at=datetime.utcnow(),
        updated_at=datetime.utcnow(),
        published_at=datetime(2024, 1, 2, 3, 4, 5),
    )
    db.add(video)
    db.commit()
    db.refresh(video)
    return video


class _YouTube:
    def __init__(self):
        self.detail_calls = []

    def get_video_details(self, video_id):
        raise AssertionError(f"unexpected single lookup for {video_id}")

    def get_videos_details(self, video_ids):
        raise AssertionError(f"unexpected batch lookup for {video_ids}")


def test_normalize_tags_keeps_a_json_array_shape():
    assert normalize_tags(None) == []
    assert normalize_tags("[]") == []
    assert normalize_tags("null") == []
    assert normalize_tags('["euv", " chips ", "euv"]') == ["euv", "chips"]
    assert normalize_tags("alpha, beta, alpha") == ["alpha", "beta"]
    assert normalize_tags(["Veritasium", "veritasium"]) == ["Veritasium", "veritasium"]
    assert tags_missing([]) is True
    assert tags_missing(["a"]) is False


def test_resolve_tags_prefers_api_and_falls_back_to_ytdlp(monkeypatch):
    def explode(video_id):
        raise AssertionError("yt-dlp should not run when the API returned tags")

    monkeypatch.setattr("app.services.video_tags.fetch_tags_via_ytdlp", explode)
    api = resolve_tags("abc", [" Owner ", "Owner", ""])
    assert api.tags == ["Owner"]
    assert api.ok is True
    assert api.source == "api"

    monkeypatch.setattr(
        "app.services.video_tags.fetch_tags_via_ytdlp",
        lambda video_id: (["public"], True),
    )
    public = resolve_tags("abc", [])
    assert public.tags == ["public"]
    assert public.source == "ytdlp"

    monkeypatch.setattr(
        "app.services.video_tags.fetch_tags_via_ytdlp",
        lambda video_id: ([], False),
    )
    failed = resolve_tags("abc", None)
    assert failed.ok is False
    assert failed.tags == []


def test_parse_video_item_reads_snippet_tags():
    service = YouTubeService.__new__(YouTubeService)
    parsed = service._parse_video_item({
        "id": "abc123",
        "snippet": {
            "title": "Chip machine",
            "tags": ["euv", "asml"],
            "publishedAt": "2024-02-05T12:00:00Z",
            "thumbnails": {"medium": {"url": "https://img.example/a.jpg"}},
        },
        "statistics": {"viewCount": "10", "likeCount": "2", "commentCount": "1"},
        "contentDetails": {"duration": "PT1M2S"},
    })
    assert parsed["video_id"] == "abc123"
    assert parsed["tags"] == ["euv", "asml"]
    assert parsed["duration"] == 62

    empty = service._parse_video_item({
        "id": "no-tags",
        "snippet": {"title": "No tags"},
        "statistics": {},
        "contentDetails": {},
    })
    assert empty["tags"] == []


def test_schema_emits_empty_list_when_tags_are_null():
    video = VideoWithChannel(
        id=1,
        video_id="dQw4w9WgXcQ",
        title="Title",
        channel_id=1,
        tags=None,
        created_at=datetime.utcnow(),
        updated_at=datetime.utcnow(),
        channel_name="Chan",
        channel_url="https://www.youtube.com/channel/UCtest",
    )
    assert video.tags == []
    assert video.model_dump()["tags"] == []


def test_process_video_stores_api_tags_and_backfills_existing(db, monkeypatch):
    monkeypatch.setattr("app.services.settings_service.SettingsService.get", staticmethod(_settings_get))
    monkeypatch.setattr("app.services.crawler_service.YouTubeService", _YouTube)

    ytdlp_calls = []

    def fake_ytdlp(video_id):
        ytdlp_calls.append(video_id)
        return ["from-ytdlp"], True

    monkeypatch.setattr("app.services.video_tags.fetch_tags_via_ytdlp", fake_ytdlp)

    channel = _channel(db)
    session = _session(db, channel)
    already = _video(db, channel, "keep-me", tags=["keep"])
    unchecked = _video(db, channel, "needs-tags", tags=[])
    checked_empty = _video(db, channel, "no-public-tags", tags=[], checked=True)

    from app.services.crawler_service import CrawlerService
    crawler = CrawlerService(db)

    details = {
        "brand-new": {
            "title": "Brand new",
            "description": "d",
            "duration": 5,
            "published_at": None,
            "view_count": 1,
            "like_count": 0,
            "comment_count": 0,
            "tags": ["api-tag"],
        },
        "needs-tags": {"tags": []},
    }

    asyncio.run(crawler.process_video(
        {"video_id": "brand-new", "title": "Brand new"},
        channel,
        session,
        details_by_id=details,
    ))
    asyncio.run(crawler.process_video(
        {"video_id": "needs-tags", "title": "Needs"},
        channel,
        session,
        details_by_id=details,
    ))
    asyncio.run(crawler.process_video(
        {"video_id": "keep-me", "title": "Keep"},
        channel,
        session,
        details_by_id={},
    ))
    asyncio.run(crawler.process_video(
        {"video_id": "no-public-tags", "title": "None"},
        channel,
        session,
        details_by_id={},
    ))

    created = db.query(Video).filter(Video.video_id == "brand-new").one()
    assert created.tags == ["api-tag"]
    assert created.tags_checked_at is not None

    db.refresh(unchecked)
    assert unchecked.tags == ["from-ytdlp"]
    assert unchecked.tags_checked_at is not None

    db.refresh(already)
    assert already.tags == ["keep"]

    db.refresh(checked_empty)
    assert checked_empty.tags == []

    assert ytdlp_calls == ["needs-tags"]


def test_backfill_writes_api_and_ytdlp_tags_without_clobbering(db, monkeypatch):
    calls = []

    def fake_ytdlp(video_id):
        calls.append(video_id)
        if video_id == "has-tags":
            return [], True
        return ["public-tag"], True

    monkeypatch.setattr("app.services.video_tags.fetch_tags_via_ytdlp", fake_ytdlp)

    channel = _channel(db, "Backfill")

    class Api:
        def get_videos_details(self, video_ids):
            return {
                video_id: {"video_id": video_id, "tags": ["api-tag"] if video_id == "from-api" else []}
                for video_id in video_ids
            }

    from_api = _video(db, channel, "from-api", tags=None)
    from_ytdlp = _video(db, channel, "from-ytdlp", tags=[])
    kept = _video(db, channel, "has-tags", tags=["keep"])

    stats = backfill_missing_tags(db, Api(), sleep_seconds=0)
    db.refresh(from_api)
    db.refresh(from_ytdlp)
    db.refresh(kept)

    assert stats["examined"] == 2
    assert stats["updated"] == 2
    assert from_api.tags == ["api-tag"]
    assert from_ytdlp.tags == ["public-tag"]
    assert kept.tags == ["keep"]
    assert calls == ["from-ytdlp"]
    assert from_api.tags_checked_at is not None

    forced = backfill_missing_tags(db, Api(), force=True, sleep_seconds=0)
    db.refresh(kept)
    assert kept.tags == ["keep"]
    assert forced["unchanged"] >= 1


def test_backfill_dry_run_does_not_write(db, monkeypatch):
    monkeypatch.setattr(
        "app.services.video_tags.fetch_tags_via_ytdlp",
        lambda video_id: (["would-write"], True),
    )

    class Api:
        def get_videos_details(self, video_ids):
            return {}

    channel = _channel(db, "Dry")
    video = _video(db, channel, "dry-id", tags=[])
    stats = backfill_missing_tags(db, Api(), sleep_seconds=0, dry_run=True)
    db.refresh(video)
    assert stats["updated"] == 1
    assert video.tags == []
    assert video.tags_checked_at is None


def test_export_appends_tags_column(db):
    channel = _channel(db, "Export")
    _video(db, channel, "yt-export", tags=["alpha", "beta"])
    video = db.query(Video).filter(Video.video_id == "yt-export").one()

    rows = list(csv.reader(io.StringIO(ExportService().export_to_csv([video]))))
    assert rows[0][0] == "Video ID"
    assert rows[0][-1] == "Tags"
    assert rows[1][0] == "yt-export"
    assert rows[1][-1] == "alpha, beta"
    # Existing columns stay put: Tags is the only addition, at the end.
    assert rows[0].index("Keywords") == len(rows[0]) - 2
    assert "Yes" == rows[1][rows[0].index("Has Summary")]

    excel = ExportService().export_to_excel([video])
    import openpyxl
    sheet = openpyxl.load_workbook(io.BytesIO(excel)).active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    values = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
    assert headers[0] == "Video ID"
    assert headers[-1] == "Tags"
    assert values[0] == "yt-export"
    assert values[-1] == "alpha, beta"


def test_list_and_detail_include_tags_array(db):
    from fastapi.encoders import jsonable_encoder

    from app.api.videos import get_video, list_videos

    channel = _channel(db, "Api")
    with_tags = _video(db, channel, "yt-with-tags", tags=["euv", "chips"])
    _video(db, channel, "yt-null-tags", tags=None)

    listing = jsonable_encoder(list_videos(
        skip=0, limit=50, channel_id=None, has_summary=None, keyword=None,
        sort_by="fetched", order="desc", db=db,
    ))
    by_youtube_id = {item["video_id"]: item for item in listing}
    assert by_youtube_id["yt-with-tags"]["tags"] == ["euv", "chips"]
    assert by_youtube_id["yt-null-tags"]["tags"] == []

    body = jsonable_encoder(get_video(with_tags.id, db))
    assert body["video_id"] == "yt-with-tags"
    assert body["tags"] == ["euv", "chips"]
    assert "tags_checked_at" not in body


def test_ensure_tag_columns_is_idempotent(tmp_path):
    from sqlalchemy import create_engine, text

    from backfill_video_tags import ensure_tag_columns

    eng = create_engine(f"sqlite:///{tmp_path / 'old.db'}")
    with eng.connect() as conn:
        conn.execute(text(
            "CREATE TABLE videos (id INTEGER PRIMARY KEY, video_id VARCHAR(255), tags JSON)"
        ))
        conn.commit()

    ensure_tag_columns(eng)
    ensure_tag_columns(eng)

    with eng.connect() as conn:
        columns = [row[1] for row in conn.execute(text("PRAGMA table_info(videos)"))]
    assert "tags_checked_at" in columns
